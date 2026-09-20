"""
Extract plain text from a user-uploaded PDF or DOCX, for use as chat context.

Decided for this use case: DOCX via **python-docx**, PDF via **PyMuPDF** --
both already dependencies (PyMuPDF for the image extraction in
document_images.py), both local and synchronous, so an attached document can
be read and used within the same chat turn. This is deliberately lighter than
the still-unbuilt "Part B" upload pipeline described in api/routes/upload.py's
docstring (LlamaParse for PDF, chunk-and-embed into clause_chunks): the text
extracted here is stuffed directly into the chat prompt for that session
rather than persisted or embedded, so there is no storage and no per-document
embedding cost.

The frontend only lets a user pick a .pdf or .docx file for chat attachment
(see ChatPage.jsx), so extract_text() does not need to gracefully handle --
or explain -- any other extension; an unsupported one is a caller bug, not a
user-facing case.
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

import docx

DOCX_EXTENSIONS = {".docx"}
PDF_EXTENSIONS = {".pdf"}

# A long attachment costs prompt tokens on every turn of the chat (the text is
# resent, unchunked, with every question -- see api/routes/chat.py). Cut it
# off rather than let one huge PDF blow the model's context budget or cost.
MAX_CHARACTERS = 50_000
TRUNCATION_NOTE = "\n\n[... document truncated at {limit:,} characters ...]"


def _extract_from_docx(data: bytes) -> str:
    document = docx.Document(BytesIO(data))

    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_from_pdf(data: bytes) -> str:
    try:
        import pymupdf
    except ImportError:
        import fitz as pymupdf  # older PyMuPDF releases

    with pymupdf.open(stream=data, filetype="pdf") as doc:
        return "\n".join(page.get_text() for page in doc)


def extract_text(data: bytes, filename: str) -> str:
    """Extract plain text from an uploaded DOCX or PDF, truncated to a safe length."""
    suffix = Path(filename).suffix.lower()
    if suffix in DOCX_EXTENSIONS:
        text = _extract_from_docx(data)
    elif suffix in PDF_EXTENSIONS:
        text = _extract_from_pdf(data)
    else:
        raise ValueError(f"unsupported document type '{suffix or filename}'")

    text = text.strip()
    if len(text) > MAX_CHARACTERS:
        text = text[:MAX_CHARACTERS] + TRUNCATION_NOTE.format(limit=MAX_CHARACTERS)
    return text

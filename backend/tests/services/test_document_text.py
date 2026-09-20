"""
Tests for app.services.document_text.extract_text.

Real DOCX/PDF bytes are built with python-docx and PyMuPDF rather than
mocked, matching tests/services/test_document_captions.py's approach --
the thing being tested is the actual library integration, so a mock would
prove nothing.
"""

from io import BytesIO

import docx
import pytest

from app.services import document_text as dt


def _docx_bytes(paragraphs=(), table_rows=None):
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table_rows:
        table = document.add_table(rows=0, cols=len(table_rows[0]))
        for row_values in table_rows:
            row = table.add_row()
            for cell, value in zip(row.cells, row_values):
                cell.text = value

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _pdf_bytes(lines):
    pymupdf = pytest.importorskip("pymupdf")

    doc = pymupdf.open()
    page = doc.new_page()
    for i, line in enumerate(lines):
        page.insert_text((72, 72 + i * 20), line, fontsize=11)
    data = doc.tobytes()
    doc.close()
    return data


def test_extracts_paragraph_text_from_docx():
    data = _docx_bytes(paragraphs=["Section 1: Site drainage", "All stormwater must be piped."])

    text = dt.extract_text(data, "spec.docx")

    assert "Section 1: Site drainage" in text
    assert "All stormwater must be piped." in text


def test_extracts_table_text_from_docx():
    data = _docx_bytes(
        paragraphs=["Fixture schedule"],
        table_rows=[["Room", "Fixture"], ["Bathroom", "Smoke alarm"]],
    )

    text = dt.extract_text(data, "spec.docx")

    assert "Room | Fixture" in text
    assert "Bathroom | Smoke alarm" in text


def test_skips_blank_paragraphs_in_docx():
    data = _docx_bytes(paragraphs=["First line", "", "   ", "Second line"])

    text = dt.extract_text(data, "spec.docx")

    assert text == "First line\nSecond line"


def test_extracts_text_from_pdf():
    data = _pdf_bytes(["Clause H1D4", "Footings must be designed to support the loads."])

    text = dt.extract_text(data, "report.pdf")

    assert "Clause H1D4" in text
    assert "Footings must be designed to support the loads." in text


def test_truncates_text_over_the_character_limit(monkeypatch):
    monkeypatch.setattr(dt, "MAX_CHARACTERS", 50)
    data = _docx_bytes(paragraphs=["x" * 200])

    text = dt.extract_text(data, "big.docx")

    assert len(text) < 200
    assert "truncated at 50 characters" in text


def test_unsupported_extension_raises_value_error():
    with pytest.raises(ValueError, match="unsupported document type"):
        dt.extract_text(b"whatever", "notes.txt")
